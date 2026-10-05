import io
import json

import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient

from app.config import Settings
from app.main import create_app


@pytest.fixture
def settings(tmp_path) -> Settings:
    return Settings(
        _env_file=None,
        data_dir=tmp_path / "data",
        enable_mock_provider=True,
        openai_api_key=None,
        anthropic_api_key=None,
        google_api_key=None,
        chunk_size=300,
        chunk_overlap=30,
        max_upload_mb=1,
    )


@pytest.fixture
def client(settings):
    with TestClient(create_app(settings)) as test_client:
        yield test_client


@pytest.fixture
def mock_kb(client) -> dict:
    response = client.post(
        "/api/knowledge-bases",
        json={"name": "Handbook", "embedding_provider": "mock", "embedding_model": "mock-hashing"},
    )
    assert response.status_code == 201, response.text
    return response.json()


def make_pdf(pages: list[str]) -> bytes:
    """A minimal valid PDF with one line of Helvetica text per page."""
    page_ids = [4 + 2 * i for i in range(len(pages))]
    kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for page_id, text in zip(page_ids, pages, strict=True):
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {page_id + 1} 0 R >>".encode()
        )
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        content = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode()
        objects.append(b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream")

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n".encode()
    return bytes(out)


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    document = DocxDocument()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def parse_sse(body: str) -> list[tuple[str, dict]]:
    events = []
    for block in body.strip().split("\n\n"):
        lines = block.splitlines()
        event = next(line[len("event: ") :] for line in lines if line.startswith("event: "))
        data = json.loads(next(line[len("data: ") :] for line in lines if line.startswith("data: ")))
        events.append((event, data))
    return events
