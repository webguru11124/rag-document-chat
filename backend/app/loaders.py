"""Turn uploaded files into text segments, keeping page numbers where a format has them."""

import codecs
import csv
import io
import re
from dataclasses import dataclass
from pathlib import PurePath

from docx import Document as DocxDocument
from pypdf import PdfReader
from pypdf.errors import PdfReadError

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".txt", ".md", ".csv")


class UnsupportedFileError(Exception):
    pass


class UnreadableFileError(Exception):
    pass


@dataclass(frozen=True)
class Segment:
    text: str
    page: int | None = None


def extension(filename: str) -> str:
    return PurePath(filename).suffix.lower()


def load(filename: str, data: bytes) -> list[Segment]:
    ext = extension(filename)
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileError(
            f"'{ext or filename}' files aren't supported. Upload {', '.join(SUPPORTED_EXTENSIONS)}."
        )
    try:
        if ext == ".pdf":
            segments = _load_pdf(data)
        elif ext == ".docx":
            segments = _load_docx(data)
        elif ext == ".csv":
            segments = _load_csv(data)
        else:
            segments = [Segment(_decode(data))]
    except (UnsupportedFileError, UnreadableFileError):
        raise
    except Exception as exc:  # parsers raise many different types on corrupt files
        raise UnreadableFileError(f"Couldn't read {filename}: {exc}") from exc

    cleaned = [Segment(_clean(s.text), s.page) for s in segments]
    return [s for s in cleaned if s.text]


def _load_pdf(data: bytes) -> list[Segment]:
    try:
        reader = PdfReader(io.BytesIO(data))
    except PdfReadError as exc:
        raise UnreadableFileError(f"This PDF looks damaged: {exc}") from exc
    if reader.is_encrypted and not reader.decrypt(""):
        raise UnreadableFileError("This PDF is password-protected.")
    return [Segment(page.extract_text() or "", number) for number, page in enumerate(reader.pages, 1)]


def _load_docx(data: bytes) -> list[Segment]:
    document = DocxDocument(io.BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            parts.append(" | ".join(cells))
    return [Segment("\n".join(parts))]


def _load_csv(data: bytes) -> list[Segment]:
    text = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    if not rows:
        return []
    header, body = rows[0], rows[1:]
    if not body:
        return [Segment(", ".join(header))]
    # One line per row, "column: value" pairs, so each chunk stays self-describing.
    lines = []
    for row in body:
        pairs = [
            f"{name.strip()}: {value.strip()}"
            for name, value in zip(header, row, strict=False)
            if value.strip()
        ]
        if pairs:
            lines.append("; ".join(pairs))
    return [Segment("\n".join(lines))]


def _decode(data: bytes) -> str:
    if data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return data.decode("utf-16")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def _clean(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t\f\v]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
